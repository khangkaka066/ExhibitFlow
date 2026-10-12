from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


TEMPLATE = Path('/Users/nguyenvokhang/Downloads/VLM/RESEARCH PROPOSAL TEMPLATE.docx')
OUT = Path('/Users/nguyenvokhang/Downloads/VLM')


def clear_document(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn('w:sectPr'):
            body.remove(child)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = 'Times New Roman'
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), 'B7B7B7')
        borders.append(el)
    tbl_pr.append(borders)
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, True, (255, 255, 255), 9)
        set_cell_shading(table.rows[0].cells[i], '1F4E78')
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value), False, None, 9)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph()
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = doc.styles['Heading 1'] if level == 1 else doc.styles['Heading 2']
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.color.rgb = RGBColor(31, 78, 121)
    r.bold = True
    r.font.size = Pt(13 if level == 1 else 11)
    return p


def add_para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.7)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.05
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        r.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        p.paragraph_format.space_after = Pt(2)
        p.add_run('•  ' + item)


def configure(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)
    for style in doc.styles:
        if hasattr(style, 'font'):
            style.font.name = 'Times New Roman'
            style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
            if style.name == 'Normal':
                style.font.size = Pt(11)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run('Research Proposal — Connectome-Inspired Sparse Visual Memory')
    run.font.name = 'Times New Roman'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(100, 100, 100)


def title_block(doc, lang):
    vi = lang == 'vi'
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('ĐỀ XUẤT NGHIÊN CỨU' if vi else 'RESEARCH PROPOSAL')
    r.bold = True; r.font.name = 'Times New Roman'; r.font.size = Pt(16)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title = ('BỘ NHỚ THỊ GIÁC THƯA LẤY CẢM HỨNG TỪ CONNECTOME CHO HIỂU VIDEO DÀI BẰNG MÔ HÌNH THỊ GIÁC–NGÔN NGỮ'
             if vi else 'CONNECTOME-INSPIRED SPARSE VISUAL MEMORY FOR LONG-VIDEO VISION–LANGUAGE UNDERSTANDING')
    r = p.add_run(title); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor(31, 78, 121)
    fields = [
        ('Tiểu ban' if vi else 'Sub-committee', 'Công nghệ Thông tin / Information Technology'),
        ('Tên nhóm' if vi else 'Group name', '[Điền tên nhóm / Insert group name]'),
        ('Tác giả' if vi else 'Authors', '[Điền họ tên thành viên / Insert team members]'),
        ('Người hướng dẫn' if vi else 'Mentor', '[Điền họ tên người hướng dẫn / Insert mentor]'),
    ]
    table = doc.add_table(rows=len(fields), cols=2); table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(fields):
        set_cell_text(table.cell(i, 0), k, True, (255,255,255), 10); set_cell_shading(table.cell(i,0), '1F4E78')
        set_cell_text(table.cell(i, 1), v, False, None, 10)
    doc.add_paragraph()


VI = {
'abstract': "Video dài tạo ra lượng lớn token thị giác, làm tăng chi phí bộ nhớ và tính toán, trong khi lấy mẫu thưa có thể bỏ mất sự kiện cần cho suy luận. Nghiên cứu này đề xuất bộ nhớ thị giác hồi quy thưa lấy cảm hứng từ connectome ruồi giấm (CISVM), dùng cấu trúc mô-đun, kết nối có hướng, hồi tiếp và nút trung tâm để cập nhật trạng thái trực tuyến và nén chuỗi đặc trưng video thành số lượng memory token cố định. Câu hỏi nghiên cứu là: với cùng ngân sách tham số, FLOPs và memory token, CISVM có bảo toàn thông tin cần thiết cho suy luận thị giác–ngôn ngữ tốt hơn các cơ chế pooling, recurrent dense, random-sparse và memory bank hiện có hay không? Hệ thống giữ cố định visual encoder, projector và language model; chỉ thay module bộ nhớ. Mô hình được đánh giá trên EgoSchema, Video-MME và một bộ kiểm tra truy xuất sự kiện có kiểm soát theo độ dài video. Chỉ số chính gồm độ chính xác câu hỏi–trả lời, độ chính xác truy xuất, hiệu năng theo khoảng cách thời gian, tỉ lệ nén, số tham số, FLOPs, VRAM và độ trễ. Nghiên cứu cũng thực hiện ablation về topology, sparsity, recurrence và cơ chế ghi/đọc. Kết quả kỳ vọng là xác định liệu cấu trúc connectome tạo ra lợi ích vượt ngoài sparsity đơn thuần, đồng thời cung cấp một module bộ nhớ gọn nhẹ, có thể gắn vào VLM mã nguồn mở.",
'keywords': 'long-video understanding; vision–language model; connectome; sparse recurrent network; visual memory; token compression',
'lit': [
"Các VLM video thường biến nhiều frame thành token rồi đưa vào language model. Khi video dài hơn, số token, VRAM và chi phí attention tăng mạnh; lấy mẫu ít frame lại làm tăng nguy cơ bỏ qua sự kiện ngắn nhưng quan trọng. LongVA mở rộng ngữ cảnh để tiếp nhận nhiều frame hơn, cho thấy hướng dense-context có hiệu quả nhưng vẫn phụ thuộc vào ngân sách context.",
"MovieChat tổ chức short-term và long-term memory để chuyển token dày thành biểu diễn thưa; MA-LMM xử lý video theo dòng và tích lũy thông tin trong memory bank. Hai công trình chứng minh rằng bộ nhớ ngoài có thể mở rộng thời lượng video mà không đưa mọi token vào LLM. Tuy nhiên, topology của bộ nhớ phần lớn được thiết kế theo quy tắc kỹ thuật hoặc pooling, chưa kiểm tra liệu cấu trúc mạng thần kinh sinh học có đem lại khả năng lưu giữ tốt hơn dưới cùng ngân sách hay không.",
"EgoSchema đo hiểu video dài khoảng ba phút bằng câu hỏi trắc nghiệm, còn Video-MME bao phủ nhiều thời lượng và kiểu nội dung. Chúng phù hợp để đánh giá suy luận tổng quát, nhưng điểm tổng hợp có thể che khuất lỗi quên thông tin ở đầu video. Vì vậy nghiên cứu bổ sung phép thử needle/event retrieval, chia kết quả theo khoảng cách giữa sự kiện và câu hỏi.",
"FlyWire cung cấp connectome não ruồi trưởng thành ở độ phân giải synapse. Các phân tích cho thấy cấu trúc có hướng, thưa, mô-đun, phân bố bậc không đồng đều và một thành phần liên thông mạnh lớn. Nghiên cứu này không tuyên bố mô phỏng sinh học đầy đủ; thay vào đó, các đặc tính topology được trích xuất và thu nhỏ thành inductive bias cho bộ nhớ học sâu. Khoảng trống chính là thiếu thí nghiệm đối chứng chứng minh topology connectome tốt hơn dense và random-sparse trong long-video VLM."
],
'necessity': "Khả năng hiểu video dài cần thiết cho trợ lý video, giáo dục, giám sát, phân tích thể thao và tra cứu kho lưu trữ. Giải pháp đưa nhiều frame vào context làm tăng chi phí; giải pháp nén mạnh có thể phá hủy thứ tự, danh tính hoặc sự kiện hiếm. Do đó cần một cơ chế nén trực tuyến có dung lượng cố định nhưng vẫn duy trì nhiều thang thời gian. Tính mới của đề tài nằm ở việc chuyển các đặc tính topology đo được từ connectome thành kiến trúc bộ nhớ có thể học, và đánh giá bằng đối chứng matched-budget để tách lợi ích của connectome khỏi lợi ích của sparsity.",
'rq': "RQ3: Liệu bộ nhớ thị giác thưa lấy cảm hứng từ connectome có thể nén thông tin video dài hiệu quả hơn trong khi vẫn bảo toàn thông tin cần thiết cho suy luận thị giác–ngôn ngữ hay không?",
'hyp': "H1: Với cùng số memory token và ngân sách tính toán, CISVM đạt độ chính xác QA và retrieval cao hơn random-sparse và pooling. H2: Mức suy giảm hiệu năng của CISVM theo độ dài video và khoảng cách truy xuất thấp hơn dense GRU/LSTM có số tham số tương đương. H3: Lợi ích mất đi hoặc giảm đáng kể khi phá vỡ cấu trúc mô-đun, hub hoặc đường hồi tiếp của topology.",
'objectives': ["Thiết kế quy trình ánh xạ thống kê connectome sang một recurrent memory graph có kích thước thực dụng.", "Tích hợp CISVM vào một VLM video mã nguồn mở trong khi giữ cố định visual encoder và LLM.", "So sánh công bằng với mean/attention pooling, dense GRU/LSTM, random-sparse recurrent memory và memory-bank baseline.", "Đo đồng thời chất lượng suy luận, khả năng nhớ theo thời gian và hiệu quả tham số/tính toán.", "Thực hiện ablation để xác định thành phần topology nào tạo ra cải thiện."],
'scope': "Đề tài tập trung vào nén đặc trưng thị giác đã được visual encoder trích xuất, không huấn luyện VLM từ đầu và không mô phỏng động lực học sinh học ở cấp neuron. Video được xử lý theo đoạn/frame; đầu ra bộ nhớ là số token cố định đưa vào LLM. Phạm vi đánh giá chính gồm EgoSchema, Video-MME và tập retrieval có kiểm soát. Ngôn ngữ đánh giá chính là tiếng Anh do benchmark; có thể bổ sung một tập nhỏ tiếng Việt nếu nguồn lực cho phép.",
'method': [
"Kiến trúc đề xuất. Visual encoder đóng băng tạo token x_t. CISVM gồm các nhóm trạng thái tương ứng các mô-đun; cạnh có hướng được tạo từ một topology connectome thu nhỏ, giữ phân bố bậc, mô-đun và tỉ lệ hồi tiếp. Cổng ghi quyết định thông tin mới; cập nhật message-passing thưa tạo h_t; readout attention sinh K memory token cố định cho LLM.",
"Thiết kế đối chứng. Tất cả biến thể dùng cùng frame sampler, visual encoder, projector, LLM, dữ liệu và lịch huấn luyện. Các baseline gồm: mean pooling; attention pooling; dense GRU/LSTM; random-sparse graph khớp số node/cạnh/bậc; topology connectome bị xáo trộn; MovieChat/MA-LMM-style memory nếu mã nguồn và tài nguyên tương thích. Báo cáo hai chế độ: matched parameters và matched FLOPs.",
"Dữ liệu và cỡ mẫu. EgoSchema có hơn 5.000 cặp hỏi–đáp trên hơn 250 giờ video; sử dụng split chính thức và không dùng test label để tinh chỉnh. Video-MME được dùng theo giao thức công bố, báo cáo theo short/medium/long. Tập synthetic/controlled retrieval được sinh từ video hợp lệ bằng cách chèn hoặc đánh dấu sự kiện mục tiêu ở các vị trí khác nhau; ít nhất 1.000 truy vấn cho mỗi dải khoảng cách để có khoảng tin cậy ổn định.",
"Chỉ số và phân tích. Chỉ số chính: QA accuracy và retrieval accuracy. Chỉ số trí nhớ: accuracy theo temporal distance, area under memory-retention curve và độ dài tại ngưỡng 90% hiệu năng ngắn hạn. Chỉ số hiệu quả: compression ratio, trainable parameters, FLOPs/video, peak VRAM, latency và throughput. Báo cáo trung bình của ba seed; bootstrap 95% CI và paired permutation/McNemar test trên cùng câu hỏi; điều chỉnh kiểm định nhiều lần khi cần.",
"Tiêu chí thành công. H1 được ủng hộ khi CISVM cải thiện ít nhất 2 điểm phần trăm accuracy so với random-sparse tốt nhất, khoảng tin cậy 95% của chênh lệch không chứa 0, trong khi FLOPs và VRAM không vượt quá 110%. Đồng thời phải vượt ít nhất một dense recurrent baseline trong chế độ matched parameters."
],
'feas': "Đề tài khả thi vì sử dụng encoder/LLM tiền huấn luyện và chỉ huấn luyện memory module cùng projector. Có thể bắt đầu bằng feature cache để giảm chi phí GPU. Mức tối thiểu khả thi chạy trên một GPU 24–48 GB với batch nhỏ, gradient accumulation và mixed precision. Chi phí chính gồm GPU, lưu trữ video và thời gian gán nhãn tập kiểm soát. Ước tính 400–800 GPU-hours cho toàn bộ baseline, ablation và ba seed; có thể giảm bằng thử nghiệm pilot và successive halving.",
'risks': ["Topology đầy đủ quá lớn: dùng coarse-graining theo cell type/module và bảo toàn thống kê graph thay vì sao chép từng neuron.", "Cải thiện đến từ số tham số: khớp parameters/FLOPs và báo cáo Pareto quality–cost.", "Rò rỉ benchmark hoặc đánh giá bằng LLM không ổn định: ưu tiên multiple-choice/exact match và split chính thức.", "Video I/O trở thành nút thắt: cache visual features, khóa sampling policy và đo riêng chi phí encoder/memory/LLM.", "Không vượt baseline: kết quả âm vẫn có giá trị nếu ablation xác định topology nào không chuyển giao; thu hẹp claim về efficiency hoặc retention."],
'expected': "Nghiên cứu kỳ vọng tạo ra: (1) một định nghĩa vận hành được của connectome-inspired sparse visual memory; (2) mã nguồn và cấu hình thí nghiệm tái lập; (3) bảng so sánh quality–efficiency trên benchmark dài; (4) memory-retention curves theo khoảng cách; (5) kết luận có kiểm soát về việc lợi ích đến từ topology connectome hay chỉ từ sparsity. Nếu thành công, module có thể được dùng cho trợ lý video, tóm tắt bài giảng và truy xuất sự kiện dài với chi phí cố định theo thời gian ở tầng memory.",
'commit': "Nhóm tác giả cam kết chịu trách nhiệm về tính chính xác, minh bạch và nguyên bản của nghiên cứu. Công cụ AI chỉ được sử dụng theo quy định của đơn vị, phải được kiểm tra bởi tác giả và được công bố trong bản cuối. AI không được dùng để tạo dữ liệu giả, bịa đặt trích dẫn hoặc thay thế việc xác minh kết quả thực nghiệm."
}

EN = {
'abstract': "Long videos generate large numbers of visual tokens, increasing memory and computation, while aggressive frame sampling may discard events required for reasoning. This study proposes a Connectome-Inspired Sparse Visual Memory (CISVM): a sparse recurrent module whose modular, directed, recurrent and hub-structured topology is derived from statistics of the fruit-fly connectome. CISVM processes frame-level features online and compresses them into a fixed number of memory tokens. The research question is whether, under matched parameter, FLOP and memory-token budgets, CISVM preserves information needed for vision–language reasoning better than pooling, dense recurrent, random-sparse and established memory-bank alternatives. The visual encoder, projector and language model are held constant; only the memory mechanism changes. Evaluation uses EgoSchema, Video-MME and a controlled long-range event-retrieval suite. Primary outcomes are question-answering accuracy, retrieval accuracy and retention over temporal distance; efficiency outcomes are compression ratio, parameters, FLOPs, peak VRAM and latency. Ablations isolate topology, sparsity, recurrence and read/write mechanisms. The expected contribution is evidence on whether connectome topology provides value beyond generic sparsity, together with a compact memory module that can be attached to an open-source video VLM.",
'keywords': 'long-video understanding; vision–language model; connectome; sparse recurrent network; visual memory; token compression',
'lit': [
"Video VLMs commonly convert sampled frames into visual tokens consumed by a language model. As videos grow longer, token count, VRAM and attention cost increase; sparse sampling can miss brief but decisive events. LongVA demonstrates the effectiveness of extending language-model context and densely sampling more frames, but this approach remains constrained by the context budget.",
"MovieChat introduces short- and long-term memory to transform dense tokens into sparse memory, while MA-LMM processes videos online and accumulates historical information in a memory bank. These works establish memory as a practical route to long-video reasoning without passing every token to the LLM. However, their memory topology is primarily engineering-designed or pooling-based; they do not test whether a biologically grounded graph improves retention under a matched budget.",
"EgoSchema evaluates three-minute long-form video understanding with multiple-choice questions, whereas Video-MME spans diverse durations and domains. Both support general evaluation, but aggregate scores may conceal forgetting of early events. This proposal therefore adds controlled needle/event retrieval and stratifies results by the temporal distance between evidence and query.",
"FlyWire provides a synapse-resolution connectome of the adult fruit-fly brain. Network analyses report directed sparse structure, heterogeneous degree, modular organization and a large strongly connected component. This project does not claim biological simulation. It extracts and scales these topological properties as an inductive bias for a trainable memory. The key gap is the absence of controlled evidence that connectome topology outperforms dense and degree-matched random sparse memories in long-video VLMs."
],
'necessity': "Long-video understanding is important for video assistants, education, surveillance, sports analysis and archive search. Passing many frames into context is expensive, whereas strong compression can destroy order, identity or rare events. A fixed-capacity online memory capable of retaining information at multiple timescales is therefore needed. The proposal is original in translating measured connectome properties into a trainable memory architecture and using matched-budget controls to separate the effect of topology from sparsity alone.",
'rq': "RQ3: Can a fruit-fly-connectome-inspired sparse visual memory efficiently compress long-video information while preserving the evidence required for vision–language reasoning?",
'hyp': "H1: Under the same memory-token and compute budgets, CISVM achieves higher QA and retrieval accuracy than random-sparse and pooling baselines. H2: Its performance degrades more slowly with video length and retrieval distance than parameter-matched dense GRU/LSTM memories. H3: the advantage decreases when modularity, hubs or recurrent paths are disrupted.",
'objectives': ["Design a reproducible mapping from connectome statistics to a practical recurrent memory graph.", "Integrate CISVM into an open-source video VLM while freezing the visual encoder and LLM.", "Compare against mean/attention pooling, dense GRU/LSTM, random-sparse recurrent memory and memory-bank baselines.", "Jointly measure reasoning quality, temporal retention and parameter/compute efficiency.", "Use ablation studies to identify which topological properties contribute to performance."],
'scope': "The study concerns compression of features produced by a pretrained visual encoder. It does not train a foundation model from scratch and does not simulate neuron-level biology. Videos are processed by frame or clip; the memory produces a fixed number of tokens for the LLM. Core evaluation covers EgoSchema, Video-MME and a controlled retrieval suite. English is the primary evaluation language because of benchmark availability; a small Vietnamese diagnostic set may be added if resources permit.",
'method': [
"Proposed architecture. A frozen visual encoder produces tokens x_t. CISVM contains groups of state units corresponding to graph modules. Directed edges are generated from a scaled connectome topology that preserves degree distribution, modularity and recurrent-path statistics. A write gate filters incoming evidence, sparse message passing updates h_t, and an attention readout emits K fixed memory tokens to the LLM.",
"Controlled design. All variants use the same frame sampler, visual encoder, projector, LLM, data and training schedule. Baselines are mean pooling, attention pooling, dense GRU/LSTM, a random sparse graph matched for nodes/edges/degrees, shuffled connectome topology, and MovieChat/MA-LMM-style memory where code and resources permit. Results are reported under matched-parameter and matched-FLOP settings.",
"Data and sample size. EgoSchema contains over 5,000 human-curated QA pairs from more than 250 hours of video; official splits will be used without tuning on test labels. Video-MME follows its published protocol with short/medium/long strata. A controlled retrieval set is constructed from eligible videos by placing or identifying target events at varied positions; at least 1,000 queries per distance band are targeted for stable confidence intervals.",
"Metrics and analysis. Primary metrics are QA accuracy and retrieval accuracy. Memory metrics are accuracy by temporal distance, area under the retention curve and the distance at which 90% of short-range performance remains. Efficiency metrics include compression ratio, trainable parameters, FLOPs/video, peak VRAM, latency and throughput. Three random seeds are reported, with bootstrap 95% confidence intervals and paired permutation or McNemar tests on common questions; multiple comparisons are corrected where applicable.",
"Success criterion. H1 is supported if CISVM improves accuracy by at least two percentage points over the strongest random-sparse control, the 95% CI for the paired difference excludes zero, and FLOPs and VRAM remain within 110%. It must also outperform at least one dense recurrent baseline under matched parameters."
],
'feas': "The project is feasible because it reuses pretrained encoders and LLMs and trains only the memory module and projector. Cached visual features reduce GPU demand. A minimum viable study can run on one 24–48 GB GPU using small batches, gradient accumulation and mixed precision. Main costs are GPU time, video storage and controlled-set annotation. The full study is estimated at 400–800 GPU-hours for baselines, ablations and three seeds, reducible through pilots and successive halving.",
'risks': ["The full topology is too large: coarse-grain by cell type/module while preserving graph statistics.", "Gains arise from parameter count: match parameters/FLOPs and report the quality–cost Pareto frontier.", "Benchmark leakage or unstable LLM judging: prioritize official multiple-choice/exact-match protocols.", "Video I/O dominates runtime: cache visual features, fix sampling and separately profile encoder, memory and LLM.", "No gain over baselines: a negative result remains informative if ablations identify which properties fail to transfer; narrow claims to efficiency or retention where justified."],
'expected': "Expected outputs are: (1) an operational definition of connectome-inspired sparse visual memory; (2) reproducible code and configurations; (3) quality–efficiency comparisons on long-video benchmarks; (4) temporal memory-retention curves; and (5) controlled conclusions about whether benefits arise from connectome topology or generic sparsity. If successful, the module could support video assistants, lecture summarization and long-horizon event retrieval with fixed memory capacity.",
'commit': "The authors accept responsibility for the accuracy, transparency and originality of this work. AI tools may be used only as permitted by institutional policy, must be verified by the authors and disclosed in the final submission. AI shall not be used to fabricate data or citations or to replace verification of experimental results."
}


REFERENCES = [
"[1] S. Dorkenwald et al., “Neuronal wiring diagram of an adult brain,” Nature, vol. 634, pp. 124–138, 2024, doi: 10.1038/s41586-024-07558-y.",
"[2] A. Lin et al., “Network statistics of the whole-brain connectome of Drosophila,” Nature, vol. 634, pp. 153–165, 2024, doi: 10.1038/s41586-024-07968-y.",
"[3] P. Schlegel et al., “Whole-brain annotation and multi-connectome cell typing of Drosophila,” Nature, vol. 634, pp. 139–152, 2024, doi: 10.1038/s41586-024-07686-5.",
"[4] E. Song et al., “MovieChat: From Dense Token to Sparse Memory for Long Video Understanding,” in Proc. IEEE/CVF CVPR, 2024.",
"[5] B. He et al., “MA-LMM: Memory-Augmented Large Multimodal Model for Long-Term Video Understanding,” in Proc. IEEE/CVF CVPR, 2024.",
"[6] K. Mangalam, R. Akshulakov, and J. Malik, “EgoSchema: A Diagnostic Benchmark for Very Long-form Video Language Understanding,” in Advances in Neural Information Processing Systems, 2023.",
"[7] C. Fu et al., “Video-MME: The First-Ever Comprehensive Evaluation Benchmark of Multi-modal LLMs in Video Analysis,” arXiv:2405.21075, 2024.",
"[8] P. Zhang et al., “Long Context Transfer from Language to Vision,” arXiv:2406.16852, 2024.",
"[9] A. Gu and T. Dao, “Mamba: Linear-Time Sequence Modeling with Selective State Spaces,” arXiv:2312.00752, 2023.",
"[10] S. Hochreiter and J. Schmidhuber, “Long Short-Term Memory,” Neural Computation, vol. 9, no. 8, pp. 1735–1780, 1997.",
]


def build(lang):
    c = VI if lang == 'vi' else EN
    vi = lang == 'vi'
    doc = Document(TEMPLATE)
    clear_document(doc); configure(doc); title_block(doc, lang)

    add_heading(doc, 'Tóm tắt' if vi else 'Abstract')
    add_para(doc, c['abstract'])
    add_para(doc, ('Từ khóa: ' if vi else 'Keywords: ') + c['keywords'], 'Từ khóa: ' if vi else 'Keywords: ')

    add_heading(doc, '1. Giới thiệu' if vi else '1. Introduction')
    add_heading(doc, '1.1. Tổng quan tài liệu' if vi else '1.1. Literature review', 2)
    for x in c['lit']: add_para(doc, x)
    add_heading(doc, '1.2. Sự cần thiết của nghiên cứu' if vi else '1.2. Necessity of the research', 2)
    add_para(doc, c['necessity'])
    add_heading(doc, '1.3. Tính khả thi' if vi else '1.3. Feasibility', 2)
    add_para(doc, c['feas'])

    add_heading(doc, '1.4. Câu hỏi và giả thuyết nghiên cứu' if vi else '1.4. Research question and hypotheses', 2)
    add_para(doc, c['rq'], 'RQ3:')
    add_para(doc, c['hyp'])

    add_heading(doc, '2. Mục tiêu nghiên cứu' if vi else '2. Research objectives')
    add_bullets(doc, c['objectives'])
    add_heading(doc, '3. Phạm vi nghiên cứu' if vi else '3. Research scope')
    add_para(doc, c['scope'])

    add_heading(doc, '4. Cách tiếp cận và phương pháp' if vi else '4. Approach and method')
    for x in c['method']: add_para(doc, x)

    headers = (['Biến thể', 'Topology', 'Ngân sách khớp', 'Mục đích'] if vi else ['Variant', 'Topology', 'Matched budget', 'Purpose'])
    rows = [
        ['Mean/attention pooling', 'None', 'Memory tokens', 'Compression baseline'],
        ['Dense GRU/LSTM', 'Dense recurrent', 'Parameters/FLOPs', 'Conventional recurrence'],
        ['Random-sparse', 'Degree/edge matched', 'Parameters/FLOPs', 'Effect of sparsity'],
        ['Shuffled-connectome', 'Perturbed modules/hubs', 'All', 'Topology ablation'],
        ['CISVM', 'Connectome-inspired', 'All', 'Proposed method'],
    ]
    add_table(doc, headers, rows, [3.2, 3.5, 3.2, 6.0])

    add_heading(doc, '4.1. Dữ liệu, chỉ số và nguồn lực' if vi else '4.1. Data, metrics and resources', 2)
    metric_headers = ['Nhóm' if vi else 'Category', 'Chỉ số' if vi else 'Measures']
    metric_rows = [
        ['Task quality', 'QA accuracy; retrieval accuracy; accuracy by duration/domain'],
        ['Memory', 'Retention curve; temporal-distance accuracy; 90% retention length'],
        ['Efficiency', 'Compression ratio; parameters; FLOPs; peak VRAM; latency; throughput'],
        ['Reliability', '3 seeds; bootstrap 95% CI; paired significance test'],
    ]
    add_table(doc, metric_headers, metric_rows, [3.7, 12.2])
    budget_headers = ['Hạng mục' if vi else 'Item', 'Ước tính' if vi else 'Estimate', 'Ghi chú' if vi else 'Notes']
    budget_rows = [
        ['GPU compute', '400–800 GPU-hours', 'Pilot, baselines, ablations, 3 seeds'],
        ['Storage', '1–3 TB', 'Video or cached visual features'],
        ['Annotation', '40–80 person-hours', 'Controlled retrieval validation'],
        ['Software/data', 'Open source', 'Subject to dataset licences'],
    ]
    add_table(doc, budget_headers, budget_rows, [4.0, 3.5, 8.5])

    add_heading(doc, '4.2. Thách thức và quản trị rủi ro' if vi else '4.2. Challenges and risk management', 2)
    add_bullets(doc, c['risks'])

    add_heading(doc, '5. Kế hoạch nghiên cứu' if vi else '5. Research plan')
    plan_headers = ['STT' if vi else 'No.', 'Thời gian' if vi else 'Period', 'Nhiệm vụ' if vi else 'Task', 'Sản phẩm' if vi else 'Output', 'Phụ trách' if vi else 'Owner']
    plan_vi = [
        ['1','Tuần 1','Rà soát tài liệu; chốt topology, dữ liệu và protocol','Đặc tả nghiên cứu','Cả nhóm'],
        ['2','Tuần 2','Chuẩn hóa dữ liệu; dựng feature cache và baseline','Data pipeline + baseline','Thành viên 1'],
        ['3','Tuần 3','Cài đặt CISVM; cân bằng tham số và FLOPs','Mã nguồn + unit tests','Thành viên 1–2'],
        ['4','Tuần 4','Huấn luyện pilot; điều chỉnh cấu hình hợp lệ','Checkpoint + cấu hình','Cả nhóm'],
        ['5','Tuần 5','Chạy benchmark, ablation và profiling','Logs + kết quả thô','Thành viên 2–3'],
        ['6','Tuần 6','Phân tích thống kê, viết báo cáo và kiểm tra tái lập','Báo cáo + repository','Cả nhóm'],
    ]
    plan_en = [
        ['1','Week 1','Review literature; finalize topology, data and protocol','Research specification','All'],
        ['2','Week 2','Prepare data, feature cache and baselines','Data pipeline + baselines','Member 1'],
        ['3','Week 3','Implement CISVM; match parameter and FLOP budgets','Code + unit tests','Members 1–2'],
        ['4','Week 4','Run pilot training and validate configurations','Checkpoint + configurations','All'],
        ['5','Week 5','Run benchmarks, ablations and profiling','Logs + raw results','Members 2–3'],
        ['6','Week 6','Statistical analysis, writing and reproducibility audit','Report + repository','All'],
    ]
    add_table(doc, plan_headers, plan_vi if vi else plan_en, [1.0, 2.3, 5.0, 4.3, 2.5])

    add_heading(doc, '6. Kết quả dự kiến' if vi else '6. Expected results')
    add_para(doc, c['expected'])

    add_heading(doc, 'Tài liệu tham khảo' if vi else 'References')
    for ref in REFERENCES:
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.7); p.paragraph_format.first_line_indent = Cm(-0.7); p.paragraph_format.space_after = Pt(3); p.add_run(ref)

    add_heading(doc, 'Phụ lục A — Ma trận ablation' if vi else 'Appendix A — Ablation matrix')
    abl_headers = ['Thí nghiệm' if vi else 'Experiment', 'Thay đổi' if vi else 'Change', 'Kết luận mục tiêu' if vi else 'Target inference']
    abl_rows = [
        ['A1', 'Connectome vs degree-matched random', 'Topology beyond sparsity'],
        ['A2', 'Remove recurrent edges', 'Value of recurrence'],
        ['A3', 'Destroy module assignments', 'Value of modularity'],
        ['A4', 'Cap high-degree hubs', 'Value of hub structure'],
        ['A5', 'Vary K and sparsity', 'Quality–efficiency frontier'],
        ['A6', 'Query-aware vs query-agnostic readout', 'Role of question conditioning'],
    ]
    add_table(doc, abl_headers, abl_rows, [2.0, 6.0, 8.0])

    add_heading(doc, 'CAM KẾT SỬ DỤNG AI CÓ TRÁCH NHIỆM' if vi else 'COMMITMENT ON RESPONSIBLE AI USAGE IN RESEARCH')
    add_para(doc, c['commit'])
    add_table(doc, ['Tác giả 1' if vi else 'Author 1', 'Tác giả 2' if vi else 'Author 2', 'Tác giả 3' if vi else 'Author 3'], [['\n\nChữ ký và họ tên' if vi else '\n\nSignature & full name']*3], [5.3,5.3,5.3])

    filename = 'De_xuat_RQ3_VLM_Connectome_Tieng_Viet.docx' if vi else 'RQ3_Connectome_VLM_Research_Proposal_English.docx'
    path = OUT / filename
    doc.save(path)
    return path


if __name__ == '__main__':
    for language in ('vi', 'en'):
        print(build(language))
